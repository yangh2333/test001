#!/usr/bin/env python3
import sys
import os
import json
import sqlite3
import smtplib
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from datetime import datetime
from collections import defaultdict

CONFIG_FILE = 'config.json'
DB_FILE = 'data/price_monitor.db'
SCHEMA_FILE = 'schema.sql'
LOG_FILE = 'logs/monitor.log'

def load_config():
    with open(CONFIG_FILE, 'r') as f:
        return json.load(f)

def init_logging():
    os.makedirs('logs', exist_ok=True)
    return LOG_FILE

def log(message):
    log_file = init_logging()
    timestamp = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    with open(log_file, 'a') as f:
        f.write(f'[{timestamp}] {message}\n')
    print(message)

def init_database():
    os.makedirs('data', exist_ok=True)
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    with open(SCHEMA_FILE, 'r') as f:
        schema = f.read()
    cursor.executescript(schema)
    conn.commit()
    conn.close()
    log('数据库初始化完成')

def get_db_connection():
    return sqlite3.connect(DB_FILE)

def get_monitor_items():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT id, category, device_name, target_upper_limit, target_lower_limit FROM monitor_items ORDER BY id')
    items = cursor.fetchall()
    conn.close()
    return items

def get_latest_price(item_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT price, record_date FROM price_records WHERE item_id = ? ORDER BY record_date DESC LIMIT 1', (item_id,))
    result = cursor.fetchone()
    conn.close()
    return result

def insert_price_record(item_id, price, source='manual'):
    conn = get_db_connection()
    cursor = conn.cursor()
    record_date = datetime.now().strftime('%Y-%m-%d')
    cursor.execute('INSERT INTO price_records (item_id, price, record_date, source) VALUES (?, ?, ?, ?)',
                  (item_id, price, record_date, source))
    conn.commit()
    conn.close()

def insert_alert(item_id, alert_type, message):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('INSERT INTO alerts (item_id, alert_type, message) VALUES (?, ?, ?)',
                  (item_id, alert_type, message))
    conn.commit()
    conn.close()

def get_unresolved_alerts():
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('SELECT a.id, a.item_id, a.alert_type, a.message, a.triggered_at, m.device_name FROM alerts a JOIN monitor_items m ON a.item_id = m.id WHERE a.is_resolved = 0 ORDER BY a.triggered_at DESC')
    alerts = cursor.fetchall()
    conn.close()
    return alerts

def resolve_alert(alert_id):
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute('UPDATE alerts SET is_resolved = 1 WHERE id = ?', (alert_id,))
    conn.commit()
    conn.close()

def send_email(subject, body):
    config = load_config()
    smtp_config = config['smtp']
    
    msg = MIMEMultipart()
    msg['From'] = smtp_config['from_address']
    msg['To'] = ', '.join(smtp_config['to_addresses'])
    msg['Subject'] = subject
    
    msg.attach(MIMEText(body, 'plain', 'utf-8'))
    
    try:
        server = smtplib.SMTP(smtp_config['server'], smtp_config['port'])
        if smtp_config.get('use_tls', True):
            server.starttls()
        server.login(smtp_config['username'], smtp_config['password'])
        text = msg.as_string()
        server.sendmail(smtp_config['from_address'], smtp_config['to_addresses'], text)
        server.quit()
        log('邮件发送成功')
        return True
    except Exception as e:
        log(f'邮件发送失败: {str(e)}')
        return False

def check_alerts(item_id, price, item):
    config = load_config()
    target_upper = item[3]
    target_lower = item[4]
    device_name = item[2]
    
    alerts = []
    
    if price > target_upper:
        message = f"价格预警【{item_id}】{device_name}: 当前价格¥{price:,.2f}超过目标上限¥{target_upper:,.2f}"
        alerts.append(('PRICE_UPPER', message))
    
    if target_lower > 0 and price < target_lower:
        message = f"价格预警【{item_id}】{device_name}: 当前价格¥{price:,.2f}低于目标下限¥{target_lower:,.2f}"
        alerts.append(('PRICE_LOWER', message))
    
    latest = get_latest_price(item_id)
    if latest:
        prev_price = latest[0]
        change = (price - prev_price) / prev_price
        increase_threshold = config['alert']['price_increase_threshold']
        decrease_threshold = config['alert']['price_decrease_threshold']
        
        if change > increase_threshold:
            message = f"涨幅预警【{item_id}】{device_name}: 价格上涨{change*100:.1f}%，¥{prev_price:,.2f} → ¥{price:,.2f}"
            alerts.append(('PRICE_INCREASE', message))
        
        if change < -decrease_threshold:
            message = f"跌幅预警【{item_id}】{device_name}: 价格下跌{abs(change)*100:.1f}%，¥{prev_price:,.2f} → ¥{price:,.2f}"
            alerts.append(('PRICE_DECREASE', message))
    
    for alert_type, message in alerts:
        log(message)
        insert_alert(item_id, alert_type, message)
    
    if alerts:
        subject = f"价格监控预警 ({len(alerts)}条)"
        body = '\n\n'.join([a[1] for a in alerts])
        send_email(subject, body)
    
    return alerts

def cmd_batch():
    init_database()
    items = get_monitor_items()
    
    print("\n===== 批量价格录入 =====")
    print("按提示逐项输入当前市场价格，输入 'q' 退出\n")
    
    for item in items:
        item_id, category, device_name, target_upper, target_lower = item
        
        latest = get_latest_price(item_id)
        latest_info = f" (上次记录: ¥{latest[0]:,.2f} [{latest[1]}])" if latest else ""
        
        while True:
            price_str = input(f"【{item_id}】{category} - {device_name}{latest_info}: ")
            if price_str.lower() == 'q':
                print("退出录入")
                return
            
            try:
                price = float(price_str)
                if price < 0:
                    print("价格不能为负数，请重新输入")
                    continue
                break
            except ValueError:
                print("无效的价格，请输入数字")
        
        insert_price_record(item_id, price)
        check_alerts(item_id, price, item)
        print(f"已记录: ¥{price:,.2f}\n")
    
    print("批量录入完成")

def cmd_report():
    items = get_monitor_items()
    records = []
    
    for item in items:
        item_id, category, device_name, target_upper, target_lower = item
        latest = get_latest_price(item_id)
        if latest:
            price, record_date = latest
            status = "正常"
            if price > target_upper:
                status = "⚠️ 超上限"
            elif target_lower > 0 and price < target_lower:
                status = "⚠️ 低于下限"
            records.append({
                'id': item_id,
                'category': category,
                'device_name': device_name,
                'price': price,
                'record_date': record_date,
                'target_upper': target_upper,
                'target_lower': target_lower,
                'status': status
            })
    
    html_content = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>价格监控报告 - {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</title>
    <style>
        body {{ font-family: Arial, sans-serif; margin: 20px; }}
        h1 {{ color: #333; text-align: center; }}
        table {{ width: 100%; border-collapse: collapse; }}
        th, td {{ padding: 12px; text-align: left; border-bottom: 1px solid #ddd; }}
        th {{ background-color: #f2f2f2; }}
        tr:hover {{ background-color: #f5f5f5; }}
        .warning {{ background-color: #ffebee; }}
        .normal {{ background-color: #e8f5e9; }}
        .footer {{ margin-top: 20px; text-align: center; color: #666; }}
    </style>
</head>
<body>
    <h1>价格监控报告</h1>
    <p style="text-align: right;">生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
    <table>
        <tr>
            <th>ID</th>
            <th>品类</th>
            <th>设备名称</th>
            <th>当前价格</th>
            <th>目标区间</th>
            <th>记录日期</th>
            <th>状态</th>
        </tr>
"""
    
    for record in records:
        status_class = "warning" if "⚠️" in record['status'] else "normal"
        html_content += f"""        <tr class="{status_class}">
            <td>{record['id']}</td>
            <td>{record['category']}</td>
            <td>{record['device_name']}</td>
            <td>¥{record['price']:,.2f}</td>
            <td>¥{record['target_lower']:,.2f} - ¥{record['target_upper']:,.2f}</td>
            <td>{record['record_date']}</td>
            <td>{record['status']}</td>
        </tr>
"""
    
    html_content += """    </table>
    <div class="footer">价格监控系统 v1.0</div>
</body>
</html>"""
    
    report_file = f"report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.html"
    with open(report_file, 'w', encoding='utf-8') as f:
        f.write(html_content)
    
    print(f"HTML报告已生成: {report_file}")

def cmd_export():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font
    except ImportError:
        print("错误: 未安装 openpyxl，请先安装: pip install openpyxl")
        sys.exit(1)
    
    items = get_monitor_items()
    conn = get_db_connection()
    cursor = conn.cursor()
    
    wb = Workbook()
    ws = wb.active
    ws.title = "价格记录"
    
    headers = ["ID", "品类", "设备名称", "价格", "记录日期", "来源", "目标上限", "目标下限"]
    ws.append(headers)
    
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col in range(1, len(headers)+1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
    
    warning_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    for item in items:
        item_id, category, device_name, target_upper, target_lower = item
        
        cursor.execute('SELECT price, record_date, source FROM price_records WHERE item_id = ? ORDER BY record_date DESC', (item_id,))
        for row in cursor.fetchall():
            price, record_date, source = row
            ws.append([item_id, category, device_name, price, record_date, source, target_upper, target_lower])
            
            row_num = ws.max_row
            if price > target_upper or (target_lower > 0 and price < target_lower):
                for col in range(1, 9):
                    ws.cell(row=row_num, column=col).fill = warning_fill
    
    conn.close()
    
    export_file = f"price_export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    wb.save(export_file)
    print(f"Excel报告已生成: {export_file}")

def cmd_check():
    alerts = get_unresolved_alerts()
    
    if not alerts:
        print("===== 当前无未处理预警 =====")
        return
    
    print("===== 当前预警列表 =====")
    print(f"共 {len(alerts)} 条未处理预警\n")
    
    for alert in alerts:
        alert_id, item_id, alert_type, message, triggered_at, device_name = alert
        print(f"【{alert_id}】{message}")
        print(f"  时间: {triggered_at}")
        print()
    
    while True:
        choice = input("输入预警ID标记为已处理，或输入 'q' 退出: ")
        if choice.lower() == 'q':
            break
        try:
            alert_id = int(choice)
            resolve_alert(alert_id)
            print(f"预警 {alert_id} 已标记为已处理")
        except ValueError:
            print("无效的输入")

def cmd_init():
    init_database()
    print("数据库初始化完成")

def cmd_help():
    print("""价格监控系统 v1.0

使用方法:
  python monitor.py <command>

命令列表:
  batch    - 批量价格录入（推荐每周一次）
  report   - 生成HTML可视化报告
  export   - 生成Excel数据报告
  check    - 查看预警列表
  init     - 初始化数据库
  help     - 显示帮助信息

示例:
  python monitor.py batch    # 按提示输入价格
  python monitor.py report   # 生成HTML报告
  python monitor.py export   # 生成Excel报告
  python monitor.py check    # 查看预警""")

def main():
    if len(sys.argv) < 2:
        cmd_help()
        sys.exit(1)
    
    command = sys.argv[1].lower()
    
    if command == 'batch':
        cmd_batch()
    elif command == 'report':
        cmd_report()
    elif command == 'export':
        cmd_export()
    elif command == 'check':
        cmd_check()
    elif command == 'init':
        cmd_init()
    elif command == 'help':
        cmd_help()
    else:
        print(f"未知命令: {command}")
        cmd_help()
        sys.exit(1)

if __name__ == '__main__':
    main()